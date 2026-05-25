from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import transaction
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.accounting import services as accounting_services
from apps.accounting.models import (
    Account,
    AccountingEvent,
    AccountingSettings,
    InventoryMovement,
    JournalEntry,
    JournalLine,
    Refund,
    RefundLine,
)
from apps.accounting.posting import queue_order_paid_accounting_event, queue_refund_completed_accounting_event
from apps.addresses.models import CustomerAddress
from apps.orders.models import Order
from apps.orders.services import add_order_item, transition_order_status
from apps.payments.models import Payment
from apps.accounting.services import (
    DEFAULT_ACCOUNT_CODES,
    DEFAULT_CHART_OF_ACCOUNTS,
    approve_refund,
    complete_refund,
    process_accounting_event,
    record_accounting_event,
    record_accounting_event_on_commit,
    register_accounting_event_handler,
    retry_accounting_event,
    seed_default_accounts_for_tenant,
)
from apps.products.models import Category, Product, ProductVariant
from apps.tenants.models import Tenant, TenantMembership


User = get_user_model()


class AccountingTestCase(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(
            name="Accounting Tenant",
            slug="accounting-tenant",
            is_active=True,
            is_default=True,
        )
        self.other_tenant = Tenant.objects.create(
            name="Other Tenant",
            slug="other-tenant",
            is_active=True,
        )
        self.user = User.objects.create_user(
            email="accounting@example.com",
            username="accounting",
            password="secret123",
        )
        TenantMembership.objects.create(
            tenant=self.tenant,
            user=self.user,
            role=TenantMembership.Role.MANAGER,
            is_active=True,
        )
        self.cash = Account.objects.create(
            tenant=self.tenant,
            code="1000",
            name="Cash",
            account_type=Account.Type.ASSET,
            normal_balance=Account.NormalBalance.DEBIT,
        )
        self.sales = Account.objects.create(
            tenant=self.tenant,
            code="4000",
            name="Sales Revenue",
            account_type=Account.Type.INCOME,
            normal_balance=Account.NormalBalance.CREDIT,
        )
        self.category = Category.objects.create(
            tenant=self.tenant,
            name="Inventory",
            slug="inventory",
        )
        self.product = Product.objects.create(
            tenant=self.tenant,
            category=self.category,
            title="Beans",
            slug="beans",
        )
        self.variant = ProductVariant.objects.create(
            tenant=self.tenant,
            product=self.product,
            name="1kg",
            sku="BEANS-1KG",
            price=Decimal("5000.00"),
            unit_cost=Decimal("3000.00"),
            stock_quantity=10,
        )

    def make_entry(self, *, debit="100.00", credit="100.00") -> JournalEntry:
        entry = JournalEntry.objects.create(
            tenant=self.tenant,
            memo="Test sale",
        )
        JournalLine.objects.create(
            tenant=self.tenant,
            journal_entry=entry,
            account=self.cash,
            debit=Decimal(debit),
        )
        JournalLine.objects.create(
            tenant=self.tenant,
            journal_entry=entry,
            account=self.sales,
            credit=Decimal(credit),
        )
        return entry


class JournalEntryModelTests(AccountingTestCase):
    def test_journal_entry_posts_when_balanced(self):
        entry = self.make_entry()

        posted = entry.post(user=self.user)

        self.assertEqual(posted.status, JournalEntry.Status.POSTED)
        self.assertEqual(posted.total_debits, Decimal("100.00"))
        self.assertEqual(posted.total_credits, Decimal("100.00"))
        self.assertIsNotNone(posted.posted_at)
        self.assertEqual(posted.posted_by, self.user)

    def test_journal_entry_cannot_post_when_unbalanced(self):
        entry = self.make_entry(debit="100.00", credit="90.00")

        with self.assertRaises(ValidationError):
            entry.post(user=self.user)

        entry.refresh_from_db()
        self.assertEqual(entry.status, JournalEntry.Status.DRAFT)

    def test_posted_journal_entry_cannot_be_edited_directly(self):
        entry = self.make_entry().post(user=self.user)

        entry.memo = "Edited after posting"
        with self.assertRaises(ValidationError):
            entry.save()

    def test_posted_journal_lines_cannot_be_edited_or_deleted(self):
        entry = self.make_entry().post(user=self.user)
        line = entry.lines.first()

        line.debit = Decimal("200.00")
        with self.assertRaises(ValidationError):
            line.save()

        with self.assertRaises(ValidationError):
            line.delete()

    def test_posted_journal_entry_cannot_be_deleted(self):
        entry = self.make_entry().post(user=self.user)

        with self.assertRaises(ValidationError):
            entry.delete()

    def test_reverse_creates_balanced_reversal_and_marks_original_reversed(self):
        entry = self.make_entry().post(user=self.user)

        reversal = entry.reverse(user=self.user, memo="Customer refund reversal")

        entry.refresh_from_db()
        self.assertEqual(entry.status, JournalEntry.Status.REVERSED)
        self.assertEqual(entry.reversed_by, self.user)
        self.assertEqual(reversal.status, JournalEntry.Status.POSTED)
        self.assertEqual(reversal.reversed_entry, entry)
        self.assertEqual(reversal.total_debits, Decimal("100.00"))
        self.assertEqual(reversal.total_credits, Decimal("100.00"))
        self.assertTrue(
            reversal.lines.filter(account=self.sales, debit=Decimal("100.00")).exists()
        )
        self.assertTrue(
            reversal.lines.filter(account=self.cash, credit=Decimal("100.00")).exists()
        )

    def test_journal_line_account_must_belong_to_same_tenant(self):
        other_account = Account.objects.create(
            tenant=self.other_tenant,
            code="1000",
            name="Other Cash",
            account_type=Account.Type.ASSET,
            normal_balance=Account.NormalBalance.DEBIT,
        )
        entry = JournalEntry.objects.create(tenant=self.tenant)

        with self.assertRaises(ValidationError):
            JournalLine.objects.create(
                tenant=self.tenant,
                journal_entry=entry,
                account=other_account,
                debit=Decimal("1.00"),
            )


class AccountingApiTests(AccountingTestCase):
    def setUp(self):
        super().setUp()
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_account_list_is_tenant_scoped(self):
        Account.objects.create(
            tenant=self.other_tenant,
            code="9999",
            name="Other Revenue",
            account_type=Account.Type.INCOME,
            normal_balance=Account.NormalBalance.CREDIT,
        )

        response = self.client.get(
            "/api/v1/accounting/accounts/",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        codes = {item["code"] for item in response.data["results"]}
        self.assertEqual(codes, {"1000", "4000"})

    def test_chart_of_accounts_create_account(self):
        response = self.client.post(
            "/api/v1/chart-of-accounts/",
            {
                "code": "1100",
                "name": "Cash on Hand",
                "account_type": "asset",
                "parent": self.cash.id,
                "description": "Till cash",
                "is_active": True,
            },
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["code"], "1100")
        self.assertEqual(response.data["account_type"], "asset")
        self.assertEqual(response.data["parent_code"], "1000")
        self.assertEqual(response.data["normal_balance"], "debit")

    def test_chart_of_accounts_root_list_is_tenant_scoped(self):
        Account.objects.create(
            tenant=self.other_tenant,
            code="9999",
            name="Other Revenue",
            account_type=Account.Type.INCOME,
            normal_balance=Account.NormalBalance.CREDIT,
        )

        response = self.client.get(
            "/api/v1/chart-of-accounts/",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        codes = {item["code"] for item in response.data["results"]}
        self.assertEqual(codes, {"1000", "4000"})

    def test_chart_of_accounts_csv_import_resolves_parent_and_creates_accounts(self):
        csv_data = (
            "code,name,account_type,parent_code,description,is_active\n"
            "2000,Liabilities,liability,,Main liability accounts,true\n"
            "2100,Accounts Payable,liability,2000,Supplier balances,true\n"
        )
        upload = SimpleUploadedFile("accounts.csv", csv_data.encode("utf-8"), content_type="text/csv")

        response = self.client.post(
            "/api/v1/chart-of-accounts/import/",
            {"file": upload},
            format="multipart",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["created_count"], 2)
        self.assertEqual(response.data["updated_count"], 0)
        payable = Account.objects.get(tenant=self.tenant, code="2100")
        self.assertEqual(payable.parent.code, "2000")

    def test_chart_of_accounts_xlsx_import_updates_when_requested(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["code", "name", "account_type", "parent_code", "description", "is_active"])
        sheet.append(["1000", "Main Cash", "asset", "", "Updated cash", False])
        sheet.append(["1200", "Bank", "asset", "1000", "Bank balances", True])
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            "accounts.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/v1/chart-of-accounts/import/",
            {"file": upload, "update_existing": "true"},
            format="multipart",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["created_count"], 1)
        self.assertEqual(response.data["updated_count"], 1)
        self.cash.refresh_from_db()
        self.assertEqual(self.cash.name, "Main Cash")
        self.assertFalse(self.cash.is_active)
        self.assertEqual(Account.objects.get(tenant=self.tenant, code="1200").parent, self.cash)

    def test_chart_of_accounts_csv_export(self):
        response = self.client.get(
            "/api/v1/chart-of-accounts/export/?format=csv",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        body = response.content.decode("utf-8")
        self.assertIn("code,name,account_type,parent_code,description,is_active", body)
        self.assertIn("1000,Cash,asset,", body)

    def test_chart_of_accounts_xlsx_export(self):
        response = self.client.get(
            "/api/v1/chart-of-accounts/export/?format=xlsx",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        sheet = workbook.active
        self.assertEqual([cell.value for cell in next(sheet.iter_rows(max_row=1))], [
            "code",
            "name",
            "account_type",
            "parent_code",
            "description",
            "is_active",
        ])

    def test_chart_of_accounts_csv_template_download(self):
        response = self.client.get(
            "/api/v1/chart-of-accounts/template/?format=csv",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertIn('filename="chart_of_accounts_template.csv"', response["Content-Disposition"])
        body = response.content.decode("utf-8")
        self.assertIn("code,name,account_type,parent_code,description,is_active", body)
        self.assertIn("1030,Payment Gateway Clearing,asset,", body)
        self.assertIn("5500,Platform and Software Fees,expense,", body)

    def test_chart_of_accounts_xlsx_template_download(self):
        response = self.client.get(
            "/api/v1/chart-of-accounts/template/?format=xlsx",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="chart_of_accounts_template.xlsx"', response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], (
            "code",
            "name",
            "account_type",
            "parent_code",
            "description",
            "is_active",
        ))
        self.assertEqual(rows[1], ("1000", "Cash", "asset", None, "Physical cash and cash-on-delivery collections.", "true"))
        self.assertEqual(rows[-1], ("5900", "General Administrative Expense", "expense", None, "Other operating and administrative expenses.", "true"))

    def test_chart_of_accounts_pdf_preview(self):
        response = self.client.get(
            "/api/v1/chart-of-accounts/preview/",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn('filename="chart-of-accounts.pdf"', response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_chart_of_accounts_import_validation_errors(self):
        csv_data = (
            "code,name,account_type,parent_code,description,is_active\n"
            ",Missing Code,asset,,,true\n"
            "8000,Invalid Type,bank,,,true\n"
            "8100,Missing Parent,expense,8999,,true\n"
        )
        upload = SimpleUploadedFile("accounts.csv", csv_data.encode("utf-8"), content_type="text/csv")

        response = self.client.post(
            "/api/v1/chart-of-accounts/import/",
            {"file": upload},
            format="multipart",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["created_count"], 0)
        self.assertEqual(response.data["skipped_count"], 3)
        self.assertEqual(len(response.data["errors"]), 3)

    def test_non_manager_cannot_access_accounting_api(self):
        staff = User.objects.create_user(
            email="staff@example.com",
            username="staff",
            password="secret123",
        )
        TenantMembership.objects.create(
            tenant=self.tenant,
            user=staff,
            role=TenantMembership.Role.STAFF,
            is_active=True,
        )
        self.client.force_authenticate(user=staff)

        response = self.client.get(
            "/api/v1/accounting/accounts/",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 403)

    def test_api_can_post_and_reverse_journal_entry(self):
        create_response = self.client.post(
            "/api/v1/accounting/journal-entries/",
            {
                "memo": "API sale",
                "lines": [
                    {"account": self.cash.id, "debit": "50.00", "credit": "0.00"},
                    {"account": self.sales.id, "debit": "0.00", "credit": "50.00"},
                ],
            },
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        self.assertEqual(create_response.status_code, 201)
        entry_id = create_response.data["id"]

        post_response = self.client.post(
            f"/api/v1/accounting/journal-entries/{entry_id}/post/",
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        self.assertEqual(post_response.status_code, 200)
        self.assertEqual(post_response.data["status"], JournalEntry.Status.POSTED)

        patch_response = self.client.patch(
            f"/api/v1/accounting/journal-entries/{entry_id}/",
            {"memo": "Edited"},
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        self.assertEqual(patch_response.status_code, 400)

        reverse_response = self.client.post(
            f"/api/v1/accounting/journal-entries/{entry_id}/reverse/",
            {"memo": "Reverse API sale"},
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        self.assertEqual(reverse_response.status_code, 201)
        self.assertEqual(reverse_response.data["status"], JournalEntry.Status.POSTED)

        original = JournalEntry.objects.get(pk=entry_id)
        self.assertEqual(original.status, JournalEntry.Status.REVERSED)

    def test_settings_endpoint_creates_tenant_settings(self):
        response = self.client.get(
            "/api/v1/accounting/settings/",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["base_currency"], "UGX")
        self.assertTrue(AccountingSettings.objects.filter(tenant=self.tenant).exists())

    def test_accounting_event_is_tenant_scoped_and_idempotent(self):
        event = AccountingEvent.objects.create(
            tenant=self.tenant,
            event_type="manual.test",
            source_model="TestSource",
            source_id="1",
            idempotency_key="event-key",
            payload={"amount": "1.00"},
        )

        response = self.client.get(
            "/api/v1/accounting/events/",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["results"][0]["id"], event.id)

        with self.assertRaises(Exception):
            AccountingEvent.objects.create(
                tenant=self.tenant,
                event_type="manual.test",
                source_model="TestSource",
                source_id="1",
                idempotency_key="event-key",
            )

    def test_accounting_event_api_create_is_idempotent(self):
        payload = {
            "event_type": "manual.api",
            "source_model": "ManualSource",
            "source_id": "A-1",
            "idempotency_key": "same-key",
            "payload": {"amount": "10.00"},
        }

        first_response = self.client.post(
            "/api/v1/accounting/events/",
            payload,
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        second_response = self.client.post(
            "/api/v1/accounting/events/",
            {**payload, "payload": {"amount": "999.00"}},
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(first_response.status_code, 201)
        self.assertEqual(second_response.status_code, 201)
        self.assertEqual(first_response.data["id"], second_response.data["id"])
        self.assertEqual(AccountingEvent.objects.count(), 1)
        self.assertEqual(AccountingEvent.objects.get().payload, {"amount": "10.00"})

    def test_retry_failed_event_endpoint(self):
        event = AccountingEvent.objects.create(
            tenant=self.tenant,
            event_type="manual.retry",
            source_model="ManualSource",
            source_id="retry-1",
            idempotency_key="retry-key",
            status=AccountingEvent.Status.FAILED,
            error_message="temporary failure",
        )
        original_handlers = accounting_services._POSTING_HANDLERS.copy()
        register_accounting_event_handler("manual.retry", lambda event: None)

        try:
            response = self.client.post(
                f"/api/v1/accounting/events/{event.id}/retry/",
                format="json",
                HTTP_X_TENANT_SLUG=self.tenant.slug,
            )
        finally:
            accounting_services._POSTING_HANDLERS.clear()
            accounting_services._POSTING_HANDLERS.update(original_handlers)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["status"], AccountingEvent.Status.PROCESSED)
        event.refresh_from_db()
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        self.assertEqual(event.error_message, "")

    def test_retry_endpoint_rejects_non_failed_event(self):
        event = AccountingEvent.objects.create(
            tenant=self.tenant,
            event_type="manual.pending",
            source_model="ManualSource",
            source_id="pending-1",
            idempotency_key="pending-key",
            status=AccountingEvent.Status.PENDING,
        )

        response = self.client.post(
            f"/api/v1/accounting/events/{event.id}/retry/",
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 400)


class AccountingEventServiceTests(AccountingTestCase):
    def test_record_accounting_event_is_idempotent(self):
        first_event, first_created = record_accounting_event(
            tenant=self.tenant,
            event_type="order.paid",
            source_model="orders.Order",
            source_id="123",
            idempotency_key="order-paid-123",
            payload={"total": "100.00"},
        )
        second_event, second_created = record_accounting_event(
            tenant=self.tenant,
            event_type="order.paid",
            source_model="orders.Order",
            source_id="123",
            idempotency_key="order-paid-123",
            payload={"total": "999.00"},
        )

        self.assertTrue(first_created)
        self.assertFalse(second_created)
        self.assertEqual(first_event.id, second_event.id)
        self.assertEqual(AccountingEvent.objects.count(), 1)
        self.assertEqual(first_event.payload, {"total": "100.00"})

    def test_record_accounting_event_on_commit_waits_for_commit(self):
        with self.captureOnCommitCallbacks(execute=True) as callbacks:
            with transaction.atomic():
                record_accounting_event_on_commit(
                    tenant=self.tenant,
                    event_type="payment.paid",
                    source_model="payments.Payment",
                    source_id="PAY-1",
                    idempotency_key="payment-paid-1",
                    payload={"amount": "50.00"},
                )
                self.assertFalse(AccountingEvent.objects.exists())

        self.assertEqual(len(callbacks), 1)
        self.assertTrue(
            AccountingEvent.objects.filter(
                tenant=self.tenant,
                event_type="payment.paid",
                source_model="payments.Payment",
                source_id="PAY-1",
                idempotency_key="payment-paid-1",
            ).exists()
        )

    def test_process_event_without_handler_records_failure(self):
        event, _created = record_accounting_event(
            tenant=self.tenant,
            event_type="unknown.event",
            source_model="Unknown",
            source_id="1",
            idempotency_key="unknown-1",
            payload={},
        )

        processed = process_accounting_event(event=event)

        self.assertEqual(processed.status, AccountingEvent.Status.FAILED)
        self.assertIn("No accounting posting handler registered", processed.error_message)
        self.assertIsNotNone(processed.processed_at)

    def test_process_event_records_handler_failure(self):
        event, _created = record_accounting_event(
            tenant=self.tenant,
            event_type="manual.fail",
            source_model="Manual",
            source_id="1",
            idempotency_key="manual-fail-1",
            payload={},
        )
        original_handlers = accounting_services._POSTING_HANDLERS.copy()

        def failing_handler(event):
            raise RuntimeError("posting failed")

        register_accounting_event_handler("manual.fail", failing_handler)
        try:
            processed = process_accounting_event(event=event)
        finally:
            accounting_services._POSTING_HANDLERS.clear()
            accounting_services._POSTING_HANDLERS.update(original_handlers)

        self.assertEqual(processed.status, AccountingEvent.Status.FAILED)
        self.assertEqual(processed.error_message, "posting failed")
        self.assertIsNotNone(processed.processed_at)

    def test_retry_failed_event_processes_with_registered_handler(self):
        event = AccountingEvent.objects.create(
            tenant=self.tenant,
            event_type="manual.retry.service",
            source_model="Manual",
            source_id="2",
            idempotency_key="manual-retry-2",
            status=AccountingEvent.Status.FAILED,
            error_message="previous failure",
        )
        original_handlers = accounting_services._POSTING_HANDLERS.copy()
        calls = []

        def successful_handler(event):
            calls.append(event.id)
            return None

        register_accounting_event_handler("manual.retry.service", successful_handler)
        try:
            retried = retry_accounting_event(event=event)
        finally:
            accounting_services._POSTING_HANDLERS.clear()
            accounting_services._POSTING_HANDLERS.update(original_handlers)

        self.assertEqual(calls, [event.id])
        self.assertEqual(retried.status, AccountingEvent.Status.PROCESSED)
        self.assertEqual(retried.error_message, "")

    def test_retry_rejects_non_failed_event(self):
        event = AccountingEvent.objects.create(
            tenant=self.tenant,
            event_type="manual.pending.service",
            source_model="Manual",
            source_id="3",
            idempotency_key="manual-pending-3",
            status=AccountingEvent.Status.PENDING,
        )

        with self.assertRaisesMessage(ValueError, "Only failed accounting events"):
            retry_accounting_event(event=event)


class InventoryMovementTests(AccountingTestCase):
    def test_inventory_movement_calculates_total_cost(self):
        movement = InventoryMovement.objects.create(
            tenant=self.tenant,
            variant=self.variant,
            movement_type=InventoryMovement.MovementType.PURCHASE,
            quantity=3,
            unit_cost=Decimal("3000.00"),
            note="Opening purchase",
        )

        self.assertEqual(movement.total_cost, Decimal("9000.00"))

    def test_inventory_movement_rejects_cross_tenant_variant(self):
        other_category = Category.objects.create(
            tenant=self.other_tenant,
            name="Other Inventory",
            slug="other-inventory",
        )
        other_product = Product.objects.create(
            tenant=self.other_tenant,
            category=other_category,
            title="Other Beans",
            slug="other-beans",
        )
        other_variant = ProductVariant.objects.create(
            tenant=self.other_tenant,
            product=other_product,
            name="1kg",
            sku="OTHER-BEANS-1KG",
            price=Decimal("5000.00"),
            unit_cost=Decimal("3000.00"),
            stock_quantity=10,
        )

        with self.assertRaises(ValidationError):
            InventoryMovement.objects.create(
                tenant=self.tenant,
                variant=other_variant,
                movement_type=InventoryMovement.MovementType.PURCHASE,
                quantity=1,
                unit_cost=Decimal("3000.00"),
            )

    def test_inventory_movement_rejects_incorrect_total_cost(self):
        movement = InventoryMovement(
            tenant=self.tenant,
            variant=self.variant,
            movement_type=InventoryMovement.MovementType.PURCHASE,
            quantity=2,
            unit_cost=Decimal("3000.00"),
            total_cost=Decimal("5000.00"),
        )

        with self.assertRaises(ValidationError):
            movement.full_clean()

    def test_inventory_movement_api_creates_movement(self):
        client = APIClient()
        client.force_authenticate(user=self.user)

        response = client.post(
            "/api/v1/accounting/inventory-movements/",
            {
                "variant": self.variant.id,
                "movement_type": InventoryMovement.MovementType.ADJUSTMENT_IN,
                "quantity": 2,
                "unit_cost": "3000.00",
                "source_model": "manual.adjustment",
                "source_id": "ADJ-1",
                "note": "Stock count correction",
            },
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["total_cost"], "6000.00")
        self.assertEqual(response.data["variant_sku"], "BEANS-1KG")
        self.assertEqual(response.data["product_title"], "Beans")
        self.assertEqual(InventoryMovement.objects.filter(tenant=self.tenant).count(), 1)

    def test_inventory_movement_api_rejects_cross_tenant_variant(self):
        other_category = Category.objects.create(
            tenant=self.other_tenant,
            name="API Other Inventory",
            slug="api-other-inventory",
        )
        other_product = Product.objects.create(
            tenant=self.other_tenant,
            category=other_category,
            title="API Other Beans",
            slug="api-other-beans",
        )
        other_variant = ProductVariant.objects.create(
            tenant=self.other_tenant,
            product=other_product,
            name="1kg",
            sku="API-OTHER-BEANS-1KG",
            price=Decimal("5000.00"),
            unit_cost=Decimal("3000.00"),
            stock_quantity=10,
        )
        client = APIClient()
        client.force_authenticate(user=self.user)

        response = client.post(
            "/api/v1/accounting/inventory-movements/",
            {
                "variant": other_variant.id,
                "movement_type": InventoryMovement.MovementType.ADJUSTMENT_IN,
                "quantity": 1,
                "unit_cost": "3000.00",
            },
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

        self.assertEqual(response.status_code, 400)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True, ENABLE_EMAIL=False)
class OrderSalesPostingTests(AccountingTestCase):
    def setUp(self):
        super().setUp()
        self.status_notifications = patch("apps.orders.signals.queue_order_status_notifications").start()
        self.order_notifications = patch("apps.orders.signals.send_order_notification").start()
        self.addCleanup(patch.stopall)
        self.address = CustomerAddress.objects.create(
            user=self.user,
            street_name="Accounting Street",
            city="Kampala",
            region=CustomerAddress.Region.KAMPALA_AREA,
        )

    def make_order(self, *, slug: str, status=Order.Status.PROCESSING, shipping="500.00", discount="300.00"):
        order = Order.objects.create(
            tenant=self.tenant,
            user=self.user,
            address=self.address,
            slug=slug,
            status=status,
            shipping_fee=Decimal(shipping),
            discount_amount=Decimal(discount),
        )
        add_order_item(
            order=order,
            variant=self.variant,
            quantity=1,
            unit_price=Decimal("2000.00"),
        )
        order.recalculate_total_price()
        if order.status != status:
            order.status = status
            order.save(update_fields=["status", "updated_at"])
        return order

    def posted_entry_for_order(self, order):
        return JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="orders.Order",
            source_id=str(order.pk),
            idempotency_key=f"order-paid-{order.pk}",
        )

    def test_cash_order_paid_posts_sales_journal(self):
        order = self.make_order(slug="cash-order")
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )

        with self.captureOnCommitCallbacks(execute=True) as callbacks:
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )

        self.assertEqual(len(callbacks), 1)
        event = AccountingEvent.objects.get(event_type="order.paid", source_id=str(order.pk))
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        entry = self.posted_entry_for_order(order)
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)
        self.assertEqual(entry.total_debits, entry.total_credits)
        self.assertTrue(entry.lines.filter(account__code="1000", debit=order.total_price).exists())
        self.assertTrue(entry.lines.filter(account__code="4000", credit=Decimal("2000.00")).exists())
        self.assertTrue(entry.lines.filter(account__code="4010", credit=Decimal("500.00")).exists())
        self.assertTrue(entry.lines.filter(account__code="4090", debit=Decimal("300.00")).exists())
        self.assertFalse(entry.lines.filter(account__code="2100").exists())

    def test_paid_order_posts_cogs_journal_with_item_metadata(self):
        order = self.make_order(slug="cogs-order", discount="0.00")
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )

        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )

        cogs_entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="orders.Order",
            source_id=str(order.pk),
            idempotency_key=f"order-paid-{order.pk}-cogs",
        )
        self.assertEqual(cogs_entry.status, JournalEntry.Status.POSTED)
        self.assertEqual(cogs_entry.total_debits, Decimal("3000.00"))
        self.assertEqual(cogs_entry.total_credits, Decimal("3000.00"))

        cogs_line = cogs_entry.lines.get(account__code="5000")
        inventory_line = cogs_entry.lines.get(account__code="1200")
        self.assertEqual(cogs_line.debit, Decimal("3000.00"))
        self.assertEqual(inventory_line.credit, Decimal("3000.00"))
        self.assertEqual(cogs_line.metadata["items"][0]["variant_sku"], "BEANS-1KG")
        self.assertEqual(cogs_line.metadata["items"][0]["unit_cost"], "3000.00")
        self.assertEqual(inventory_line.metadata["items"][0]["total_cost"], "3000.00")

    def test_cogs_uses_product_cost_when_variant_cost_snapshot_is_missing(self):
        self.product.cost_price = Decimal("2500.00")
        self.product.save(update_fields=["cost_price"])
        self.variant.unit_cost = Decimal("0.00")
        self.variant.save(update_fields=["unit_cost"])
        order = self.make_order(slug="product-cost-cogs-order", discount="0.00")
        item = order.items.get()
        item.cost_price_snapshot = Decimal("0.00")
        item.save(update_fields=["cost_price_snapshot"])
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )

        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )

        cogs_entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="orders.Order",
            source_id=str(order.pk),
            idempotency_key=f"order-paid-{order.pk}-cogs",
        )
        cogs_line = cogs_entry.lines.get(account__code="5000")
        self.assertEqual(cogs_line.debit, Decimal("2500.00"))
        self.assertEqual(cogs_line.metadata["items"][0]["cost_source"], "product_current")

    def test_missing_cost_fails_order_paid_event_without_journals(self):
        AccountingSettings.objects.create(
            tenant=self.tenant,
            metadata={"missing_cost_policy": "fail"},
        )
        self.variant.unit_cost = Decimal("0.00")
        self.variant.save(update_fields=["unit_cost"])
        order = self.make_order(slug="missing-cost-order")
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )

        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )

        event = AccountingEvent.objects.get(event_type="order.paid", source_id=str(order.pk))
        self.assertEqual(event.status, AccountingEvent.Status.FAILED)
        self.assertIn("Missing product cost", event.error_message)
        self.assertFalse(
            JournalEntry.objects.filter(source_model="orders.Order", source_id=str(order.pk)).exists()
        )

    def test_missing_cost_attention_policy_posts_sales_and_marks_event_payload(self):
        AccountingSettings.objects.create(
            tenant=self.tenant,
            metadata={"missing_cost_policy": "attention"},
        )
        self.variant.unit_cost = Decimal("0.00")
        self.variant.save(update_fields=["unit_cost"])
        order = self.make_order(slug="missing-cost-attention-order")
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )

        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )

        event = AccountingEvent.objects.get(event_type="order.paid", source_id=str(order.pk))
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        self.assertTrue(event.payload["requires_attention"])
        self.assertEqual(event.payload["missing_cost_items"][0]["variant_sku"], "BEANS-1KG")
        self.assertTrue(
            JournalEntry.objects.filter(
                source_model="orders.Order",
                source_id=str(order.pk),
                idempotency_key=f"order-paid-{order.pk}",
            ).exists()
        )
        self.assertFalse(
            JournalEntry.objects.filter(
                source_model="orders.Order",
                source_id=str(order.pk),
                idempotency_key=f"order-paid-{order.pk}-cogs",
            ).exists()
        )

    def test_online_paid_order_posts_to_mobile_money(self):
        order = self.make_order(slug="online-order", discount="0.00")
        payment = Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.MTN,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )
        Payment.objects.filter(pk=payment.pk).update(status=Payment.Status.PAID)

        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="MTN payment confirmed",
            )

        event = AccountingEvent.objects.get(event_type="order.paid", source_id=str(order.pk))
        entry = event.journal_entry
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)
        self.assertTrue(entry.lines.filter(account__code="1020", debit=order.total_price).exists())
        self.assertFalse(entry.lines.filter(account__code="4090").exists())

    def test_duplicate_order_paid_event_does_not_post_twice(self):
        order = self.make_order(slug="duplicate-order")
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )

        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )
        with self.captureOnCommitCallbacks(execute=True):
            queue_order_paid_accounting_event(order)

        self.assertEqual(AccountingEvent.objects.filter(event_type="order.paid", source_id=str(order.pk)).count(), 1)
        self.assertEqual(JournalEntry.objects.filter(source_model="orders.Order", source_id=str(order.pk)).count(), 2)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True, ENABLE_EMAIL=False)
class PaymentPostingTests(AccountingTestCase):
    def setUp(self):
        super().setUp()
        self.status_notifications = patch("apps.orders.signals.queue_order_status_notifications").start()
        self.order_notifications = patch("apps.orders.signals.send_order_notification").start()
        self.addCleanup(patch.stopall)
        self.address = CustomerAddress.objects.create(
            user=self.user,
            street_name="Payment Street",
            city="Kampala",
            region=CustomerAddress.Region.KAMPALA_AREA,
        )

    def make_order_with_payment(self, *, slug: str, provider: str):
        order = Order.objects.create(
            tenant=self.tenant,
            user=self.user,
            address=self.address,
            slug=slug,
            status=Order.Status.PROCESSING,
        )
        add_order_item(
            order=order,
            variant=self.variant,
            quantity=1,
            unit_price=Decimal("2000.00"),
        )
        order.recalculate_total_price()
        order.status = Order.Status.PROCESSING
        order.save(update_fields=["status", "updated_at"])
        payment = Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=provider,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )
        return order, payment

    def post_order_sale_to_ar(self, *, slug: str, provider: str):
        order, payment = self.make_order_with_payment(slug=slug, provider=provider)
        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Sale posted before payment settlement",
            )
        sales_entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="orders.Order",
            source_id=str(order.pk),
            idempotency_key=f"order-paid-{order.pk}",
        )
        self.assertTrue(sales_entry.lines.filter(account__code="1100", debit=order.total_price).exists())
        return order, payment

    def mark_payment_paid(self, payment):
        payment.status = Payment.Status.PAID
        with self.captureOnCommitCallbacks(execute=True):
            payment.save(update_fields=["status", "paid_at", "updated_at"])
        payment.refresh_from_db()
        return payment

    def test_payment_paid_clears_accounts_receivable_for_each_provider(self):
        cases = [
            (Payment.Provider.MTN, "1020"),
            (Payment.Provider.CARD, "1010"),
            (Payment.Provider.STRIPE, "1010"),
            (Payment.Provider.PAYSTACK, "1010"),
            (Payment.Provider.FLUTTERWAVE, "1010"),
        ]

        for index, (provider, debit_account_code) in enumerate(cases, start=1):
            with self.subTest(provider=provider):
                order, payment = self.post_order_sale_to_ar(
                    slug=f"payment-ar-{index}",
                    provider=provider,
                )

                self.mark_payment_paid(payment)

                event = AccountingEvent.objects.get(
                    event_type="payment.paid",
                    source_id=str(payment.pk),
                )
                self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
                entry = JournalEntry.objects.get(
                    tenant=self.tenant,
                    source_model="payments.Payment",
                    source_id=str(payment.pk),
                    idempotency_key=f"payment-paid-{payment.pk}",
                )
                self.assertEqual(entry.status, JournalEntry.Status.POSTED)
                self.assertEqual(entry.total_debits, entry.total_credits)
                self.assertEqual(entry.total_debits, order.total_price)
                self.assertTrue(entry.lines.filter(account__code=debit_account_code, debit=order.total_price).exists())
                self.assertTrue(entry.lines.filter(account__code="1100", credit=order.total_price).exists())

    def test_cash_payment_paid_after_cash_sale_does_not_duplicate_cash(self):
        order, payment = self.make_order_with_payment(
            slug="cash-payment-no-duplicate",
            provider=Payment.Provider.CASH,
        )
        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash sale posted directly to cash",
            )
        sales_entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="orders.Order",
            source_id=str(order.pk),
            idempotency_key=f"order-paid-{order.pk}",
        )
        self.assertTrue(sales_entry.lines.filter(account__code="1000", debit=order.total_price).exists())

        self.mark_payment_paid(payment)

        event = AccountingEvent.objects.get(event_type="payment.paid", source_id=str(payment.pk))
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        self.assertTrue(event.payload["ignored"])
        self.assertIn("did not post to accounts receivable", event.payload["ignored_reason"])
        self.assertFalse(
            JournalEntry.objects.filter(
                source_model="payments.Payment",
                source_id=str(payment.pk),
            ).exists()
        )

    def test_duplicate_paid_status_save_does_not_emit_duplicate_payment_event(self):
        _order, payment = self.post_order_sale_to_ar(
            slug="duplicate-payment-save",
            provider=Payment.Provider.MTN,
        )

        self.mark_payment_paid(payment)
        with self.captureOnCommitCallbacks(execute=True) as callbacks:
            payment.save(update_fields=["status", "updated_at"])

        self.assertEqual(len(callbacks), 0)
        self.assertEqual(AccountingEvent.objects.filter(event_type="payment.paid", source_id=str(payment.pk)).count(), 1)
        self.assertEqual(JournalEntry.objects.filter(source_model="payments.Payment", source_id=str(payment.pk)).count(), 1)

    def test_payment_success_posts_sales_to_financials_without_order_paid_status(self):
        order, payment = self.make_order_with_payment(
            slug="paid-payment-posts-sale",
            provider=Payment.Provider.MTN,
        )

        self.mark_payment_paid(payment)
        order.refresh_from_db()

        self.assertEqual(order.status, Order.Status.PROCESSING)
        sale_event = AccountingEvent.objects.get(
            event_type="order.paid",
            source_id=str(order.pk),
        )
        self.assertEqual(sale_event.status, AccountingEvent.Status.PROCESSED)
        sales_entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="orders.Order",
            source_id=str(order.pk),
            idempotency_key=f"order-paid-{order.pk}",
        )
        self.assertEqual(sales_entry.status, JournalEntry.Status.POSTED)
        self.assertTrue(sales_entry.lines.filter(account__code="1020", debit=order.total_price).exists())
        self.assertTrue(sales_entry.lines.filter(account__code="4000", credit=Decimal("2000.00")).exists())

        payment_event = AccountingEvent.objects.get(
            event_type="payment.paid",
            source_id=str(payment.pk),
        )
        self.assertEqual(payment_event.status, AccountingEvent.Status.PROCESSED)
        self.assertTrue(payment_event.payload["ignored"])
        self.assertFalse(
            JournalEntry.objects.filter(
                source_model="payments.Payment",
                source_id=str(payment.pk),
            ).exists()
        )


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True, ENABLE_EMAIL=False)
class RefundPostingTests(AccountingTestCase):
    def setUp(self):
        super().setUp()
        self.status_notifications = patch("apps.orders.signals.queue_order_status_notifications").start()
        self.order_notifications = patch("apps.orders.signals.send_order_notification").start()
        self.addCleanup(patch.stopall)
        self.address = CustomerAddress.objects.create(
            user=self.user,
            street_name="Refund Street",
            city="Kampala",
            region=CustomerAddress.Region.KAMPALA_AREA,
        )

    def make_paid_order(self, *, slug: str, quantity: int = 1):
        order = Order.objects.create(
            tenant=self.tenant,
            user=self.user,
            address=self.address,
            slug=slug,
            status=Order.Status.PROCESSING,
        )
        item = add_order_item(
            order=order,
            variant=self.variant,
            quantity=quantity,
            unit_price=Decimal("2000.00"),
        )
        order.recalculate_total_price()
        order.status = Order.Status.PROCESSING
        order.save(update_fields=["status", "updated_at"])
        payment = Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PENDING,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )
        with self.captureOnCommitCallbacks(execute=True):
            transition_order_status(
                order=order,
                new_status=Order.Status.PAID,
                changed_by=self.user,
                note="Cash collected",
            )
        return order, item, payment

    def approve_and_complete(self, refund):
        refund = approve_refund(refund=refund, user=self.user)
        with self.captureOnCommitCallbacks(execute=True):
            refund = complete_refund(refund=refund, user=self.user)
        return refund

    def test_full_refund_posts_cash_sales_return_and_inventory_reversal(self):
        order, item, payment = self.make_paid_order(slug="full-refund-order")
        starting_stock = self.variant.stock_quantity
        refund = Refund.objects.create(
            tenant=self.tenant,
            order=order,
            payment=payment,
            amount=order.total_price,
            reason="Customer returned item",
            idempotency_key="full-refund-1",
        )
        RefundLine.objects.create(
            tenant=self.tenant,
            refund=refund,
            order_item=item,
            quantity=1,
            amount=order.total_price,
            return_to_stock=True,
        )

        self.approve_and_complete(refund)

        event = AccountingEvent.objects.get(event_type="refund.completed", source_id=str(refund.pk))
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="accounting.Refund",
            source_id=str(refund.pk),
            idempotency_key=f"refund-completed-{refund.pk}",
        )
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)
        self.assertEqual(entry.total_debits, entry.total_credits)
        self.assertTrue(entry.lines.filter(account__code="4100", debit=order.total_price).exists())
        self.assertTrue(entry.lines.filter(account__code="1000", credit=order.total_price).exists())
        self.assertTrue(entry.lines.filter(account__code="1200", debit=Decimal("3000.00")).exists())
        self.assertTrue(entry.lines.filter(account__code="5000", credit=Decimal("3000.00")).exists())
        inventory_line = entry.lines.get(account__code="1200")
        self.assertEqual(inventory_line.metadata["items"][0]["variant_sku"], "BEANS-1KG")
        self.variant.refresh_from_db()
        self.assertEqual(self.variant.stock_quantity, starting_stock + 1)
        self.assertTrue(
            InventoryMovement.objects.filter(
                tenant=self.tenant,
                movement_type=InventoryMovement.MovementType.RETURN,
                source_model="accounting.Refund",
                source_id=str(refund.pk),
            ).exists()
        )

    def test_partial_refund_posts_reduced_cash_and_sales_return_only(self):
        order, _item, payment = self.make_paid_order(slug="partial-refund-order", quantity=2)
        refund = Refund.objects.create(
            tenant=self.tenant,
            order=order,
            payment=payment,
            amount=Decimal("1000.00"),
            tax_amount=Decimal("100.00"),
            reason="Partial goodwill refund",
            idempotency_key="partial-refund-1",
        )

        self.approve_and_complete(refund)

        entry = JournalEntry.objects.get(
            tenant=self.tenant,
            source_model="accounting.Refund",
            source_id=str(refund.pk),
            idempotency_key=f"refund-completed-{refund.pk}",
        )
        self.assertEqual(entry.total_debits, Decimal("1000.00"))
        self.assertEqual(entry.total_credits, Decimal("1000.00"))
        self.assertTrue(entry.lines.filter(account__code="4100", debit=Decimal("900.00")).exists())
        self.assertTrue(entry.lines.filter(account__code="2100", debit=Decimal("100.00")).exists())
        self.assertTrue(entry.lines.filter(account__code="1000", credit=Decimal("1000.00")).exists())
        self.assertFalse(entry.lines.filter(account__code="1200").exists())
        self.assertFalse(entry.lines.filter(account__code="5000").exists())

    def test_duplicate_refund_completion_does_not_post_twice(self):
        order, _item, payment = self.make_paid_order(slug="duplicate-refund-order")
        refund = Refund.objects.create(
            tenant=self.tenant,
            order=order,
            payment=payment,
            amount=Decimal("500.00"),
            reason="Duplicate prevention",
            idempotency_key="duplicate-refund-1",
        )

        self.approve_and_complete(refund)
        refund.refresh_from_db()
        with self.captureOnCommitCallbacks(execute=True):
            queue_refund_completed_accounting_event(refund)

        self.assertEqual(AccountingEvent.objects.filter(event_type="refund.completed", source_id=str(refund.pk)).count(), 1)
        self.assertEqual(JournalEntry.objects.filter(source_model="accounting.Refund", source_id=str(refund.pk)).count(), 1)

    def test_refund_api_create_approve_and_complete(self):
        order, item, payment = self.make_paid_order(slug="api-refund-order")
        client = APIClient()
        client.force_authenticate(user=self.user)

        create_response = client.post(
            "/api/v1/accounting/refunds/",
            {
                "order": order.id,
                "payment": payment.id,
                "amount": "500.00",
                "reason": "API refund",
                "idempotency_key": "api-refund-1",
                "lines": [
                    {
                        "order_item": item.id,
                        "quantity": 1,
                        "amount": "500.00",
                        "return_to_stock": False,
                    }
                ],
            },
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        self.assertEqual(create_response.status_code, 201)
        refund_id = create_response.data["id"]

        approve_response = client.post(
            f"/api/v1/accounting/refunds/{refund_id}/approve/",
            format="json",
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )
        self.assertEqual(approve_response.status_code, 200)
        self.assertEqual(approve_response.data["status"], Refund.Status.APPROVED)

        with self.captureOnCommitCallbacks(execute=True):
            complete_response = client.post(
                f"/api/v1/accounting/refunds/{refund_id}/complete/",
                format="json",
                HTTP_X_TENANT_SLUG=self.tenant.slug,
            )
        self.assertEqual(complete_response.status_code, 200)
        self.assertEqual(complete_response.data["status"], Refund.Status.COMPLETED)
        self.assertTrue(
            AccountingEvent.objects.filter(event_type="refund.completed", source_id=str(refund_id)).exists()
        )


class AccountingReportApiTests(AccountingTestCase):
    def setUp(self):
        super().setUp()
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        seed_default_accounts_for_tenant(tenant=self.tenant)
        self.cash = Account.objects.get(tenant=self.tenant, code="1000")
        self.inventory = Account.objects.get(tenant=self.tenant, code="1200")
        self.equity = Account.objects.get(tenant=self.tenant, code="3000")
        self.sales = Account.objects.get(tenant=self.tenant, code="4000")
        self.cogs = Account.objects.get(tenant=self.tenant, code="5000")
        self.other_cash = Account.objects.create(
            tenant=self.other_tenant,
            code="1000",
            name="Other Cash",
            account_type=Account.Type.ASSET,
            normal_balance=Account.NormalBalance.DEBIT,
        )
        self.other_sales = Account.objects.create(
            tenant=self.other_tenant,
            code="4000",
            name="Other Sales",
            account_type=Account.Type.INCOME,
            normal_balance=Account.NormalBalance.CREDIT,
        )
        self.create_posted_entry(
            entry_date=date(2026, 1, 1),
            memo="Opening balances",
            lines=[
                (self.cash, "1300.00", "0.00"),
                (self.inventory, "300.00", "0.00"),
                (self.equity, "0.00", "1600.00"),
            ],
        )
        self.create_posted_entry(
            entry_date=date(2026, 2, 1),
            memo="Period sale",
            lines=[
                (self.cash, "500.00", "0.00"),
                (self.sales, "0.00", "500.00"),
            ],
        )
        self.create_posted_entry(
            entry_date=date(2026, 2, 2),
            memo="Period COGS",
            lines=[
                (self.cogs, "200.00", "0.00"),
                (self.inventory, "0.00", "200.00"),
            ],
        )
        self.create_draft_entry()
        self.create_other_tenant_entry()

    def create_posted_entry(self, *, entry_date, memo, lines, tenant=None):
        tenant = tenant or self.tenant
        entry = JournalEntry.objects.create(
            tenant=tenant,
            entry_date=entry_date,
            memo=memo,
        )
        for account, debit, credit in lines:
            JournalLine.objects.create(
                tenant=tenant,
                journal_entry=entry,
                account=account,
                debit=Decimal(debit),
                credit=Decimal(credit),
            )
        return entry.post(user=self.user)

    def create_draft_entry(self):
        entry = JournalEntry.objects.create(
            tenant=self.tenant,
            entry_date=date(2026, 2, 3),
            memo="Draft ignored",
        )
        JournalLine.objects.create(
            tenant=self.tenant,
            journal_entry=entry,
            account=self.cash,
            debit=Decimal("9999.00"),
        )
        JournalLine.objects.create(
            tenant=self.tenant,
            journal_entry=entry,
            account=self.sales,
            credit=Decimal("9999.00"),
        )

    def create_other_tenant_entry(self):
        self.create_posted_entry(
            tenant=self.other_tenant,
            entry_date=date(2026, 2, 1),
            memo="Other tenant ignored",
            lines=[
                (self.other_cash, "777.00", "0.00"),
                (self.other_sales, "0.00", "777.00"),
            ],
        )

    def get_report(self, path, params=None):
        return self.client.get(
            path,
            params or {},
            HTTP_X_TENANT_SLUG=self.tenant.slug,
        )

    def test_general_ledger_report_uses_posted_lines_opening_and_tenant_scope(self):
        response = self.get_report(
            "/api/v1/accounting/reports/general-ledger/",
            {"date_from": "2026-02-01", "date_to": "2026-02-28"},
        )

        self.assertEqual(response.status_code, 200)
        cash = next(item for item in response.data["accounts"] if item["account_code"] == "1000")
        self.assertEqual(cash["opening_balance"], Decimal("1300.00"))
        self.assertEqual(cash["period_debits"], Decimal("500.00"))
        self.assertEqual(cash["period_credits"], Decimal("0.00"))
        self.assertEqual(cash["closing_balance"], Decimal("1800.00"))
        self.assertEqual(len(cash["lines"]), 1)
        self.assertNotIn("777.00", str(response.data))
        self.assertNotIn("9999.00", str(response.data))

    def test_trial_balance_report_balances_and_includes_opening(self):
        response = self.get_report(
            "/api/v1/accounting/reports/trial-balance/",
            {"date_from": "2026-02-01", "date_to": "2026-02-28"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["is_balanced"])
        self.assertEqual(response.data["totals"]["opening_debit"], Decimal("1600.00"))
        self.assertEqual(response.data["totals"]["opening_credit"], Decimal("1600.00"))
        self.assertEqual(response.data["totals"]["period_debit"], Decimal("700.00"))
        self.assertEqual(response.data["totals"]["period_credit"], Decimal("700.00"))
        self.assertEqual(response.data["totals"]["closing_debit"], Decimal("2100.00"))
        self.assertEqual(response.data["totals"]["closing_credit"], Decimal("2100.00"))

    def test_profit_and_loss_report_uses_period_income_and_expenses(self):
        response = self.get_report(
            "/api/v1/accounting/reports/profit-and-loss/",
            {"date_from": "2026-02-01", "date_to": "2026-02-28"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["total_income"], Decimal("500.00"))
        self.assertEqual(response.data["total_expenses"], Decimal("200.00"))
        self.assertEqual(response.data["net_income"], Decimal("300.00"))

    def test_balance_sheet_report_includes_current_earnings_and_balances(self):
        response = self.get_report(
            "/api/v1/accounting/reports/balance-sheet/",
            {"date_to": "2026-02-28"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["is_balanced"])
        self.assertEqual(response.data["total_assets"], Decimal("1900.00"))
        self.assertEqual(response.data["total_liabilities"], Decimal("0.00"))
        self.assertEqual(response.data["total_equity"], Decimal("1900.00"))
        current_earnings = next(item for item in response.data["equity"] if item["account_code"] == "CURRENT_EARNINGS")
        self.assertEqual(current_earnings["amount"], Decimal("300.00"))

    def test_cash_flow_report_includes_opening_and_period_cash_change(self):
        response = self.get_report(
            "/api/v1/accounting/reports/cash-flow/",
            {"date_from": "2026-02-01", "date_to": "2026-02-28"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["opening_cash"], Decimal("1300.00"))
        self.assertEqual(response.data["net_cash_change"], Decimal("500.00"))
        self.assertEqual(response.data["closing_cash"], Decimal("1800.00"))
        cash = next(item for item in response.data["cash_accounts"] if item["account_code"] == "1000")
        self.assertEqual(cash["cash_inflows"], Decimal("500.00"))
        self.assertEqual(cash["cash_outflows"], Decimal("0.00"))

    def test_report_date_serializer_rejects_invalid_date_range(self):
        response = self.get_report(
            "/api/v1/accounting/reports/profit-and-loss/",
            {"date_from": "2026-03-01", "date_to": "2026-02-01"},
        )

        self.assertEqual(response.status_code, 400)

    def test_sales_report_uses_posted_ledger_totals(self):
        response = self.get_report(
            "/api/v1/accounting/reports/sales/",
            {"date_from": "2026-02-01", "date_to": "2026-02-28"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["totals"]["gross_sales"], Decimal("500.00"))
        self.assertEqual(response.data["totals"]["net_sales"], Decimal("500.00"))
        self.assertEqual(response.data["totals"]["cost_of_goods_sold"], Decimal("200.00"))
        self.assertEqual(response.data["totals"]["gross_profit"], Decimal("300.00"))
        self.assertEqual(response.data["totals"]["gross_margin_percent"], Decimal("60.00"))
        self.assertNotIn("777.00", str(response.data))
        self.assertNotIn("9999.00", str(response.data))


class DefaultChartOfAccountsTests(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(
            name="Seed Tenant",
            slug="seed-tenant",
            is_active=True,
            currency="UGX",
        )
        self.eur_tenant = Tenant.objects.create(
            name="Euro Tenant",
            slug="euro-tenant",
            is_active=True,
            currency="EUR",
        )

    def test_seed_creates_required_default_accounts_for_tenant(self):
        result = seed_default_accounts_for_tenant(tenant=self.tenant)

        self.assertEqual(result["created_count"], len(DEFAULT_CHART_OF_ACCOUNTS))
        self.assertEqual(result["existing_count"], 0)
        self.assertEqual(
            set(Account.objects.filter(tenant=self.tenant).values_list("code", flat=True)),
            set(DEFAULT_ACCOUNT_CODES),
        )

        defaults_by_code = {
            default.code: default
            for default in DEFAULT_CHART_OF_ACCOUNTS
        }
        for account in Account.objects.filter(tenant=self.tenant):
            default = defaults_by_code[account.code]
            self.assertEqual(account.name, default.name)
            self.assertEqual(account.account_type, default.account_type)
            self.assertEqual(account.normal_balance, default.normal_balance)
            self.assertEqual(account.currency, "UGX")
            self.assertTrue(account.is_active)

    def test_seed_uses_tenant_currency(self):
        seed_default_accounts_for_tenant(tenant=self.eur_tenant)

        self.assertEqual(
            set(Account.objects.filter(tenant=self.eur_tenant).values_list("currency", flat=True)),
            {"EUR"},
        )

    def test_seed_is_idempotent_and_preserves_existing_accounts(self):
        Account.objects.create(
            tenant=self.tenant,
            code="1000",
            name="Till Cash",
            account_type=Account.Type.ASSET,
            normal_balance=Account.NormalBalance.DEBIT,
            currency="UGX",
            description="Operator customized account",
            is_active=False,
        )

        first_result = seed_default_accounts_for_tenant(tenant=self.tenant)
        second_result = seed_default_accounts_for_tenant(tenant=self.tenant)

        self.assertEqual(first_result["created_count"], len(DEFAULT_CHART_OF_ACCOUNTS) - 1)
        self.assertEqual(first_result["existing_count"], 1)
        self.assertEqual(second_result["created_count"], 0)
        self.assertEqual(second_result["existing_count"], len(DEFAULT_CHART_OF_ACCOUNTS))
        self.assertEqual(Account.objects.filter(tenant=self.tenant).count(), len(DEFAULT_CHART_OF_ACCOUNTS))

        cash = Account.objects.get(tenant=self.tenant, code="1000")
        self.assertEqual(cash.name, "Till Cash")
        self.assertEqual(cash.description, "Operator customized account")
        self.assertFalse(cash.is_active)

    def test_management_command_seeds_all_active_tenants_idempotently(self):
        inactive_tenant = Tenant.objects.create(
            name="Inactive Tenant",
            slug="inactive-tenant",
            is_active=False,
        )
        output = StringIO()

        call_command("seed_accounting_accounts", stdout=output)
        first_count = Account.objects.count()
        call_command("seed_accounting_accounts", stdout=StringIO())

        self.assertEqual(first_count, len(DEFAULT_CHART_OF_ACCOUNTS) * 2)
        self.assertEqual(Account.objects.count(), first_count)
        self.assertFalse(Account.objects.filter(tenant=inactive_tenant).exists())
        self.assertIn("seed-tenant", output.getvalue())
        self.assertIn("euro-tenant", output.getvalue())

    def test_management_command_can_seed_selected_inactive_tenant(self):
        inactive_tenant = Tenant.objects.create(
            name="Inactive Selected",
            slug="inactive-selected",
            is_active=False,
        )

        call_command(
            "seed_accounting_accounts",
            "--tenant",
            inactive_tenant.slug,
            stdout=StringIO(),
        )

        self.assertEqual(
            Account.objects.filter(tenant=inactive_tenant).count(),
            len(DEFAULT_CHART_OF_ACCOUNTS),
        )
        self.assertFalse(Account.objects.filter(tenant=self.tenant).exists())

    def test_management_command_rejects_unknown_tenant_slug(self):
        with self.assertRaisesMessage(Exception, "Tenant slug(s) not found"):
            call_command(
                "seed_accounting_accounts",
                "--tenant",
                "missing-tenant",
                stdout=StringIO(),
            )


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True, ENABLE_EMAIL=False)
class BackfillAccountingCommandTests(AccountingTestCase):
    def setUp(self):
        super().setUp()
        self.status_notifications = patch("apps.orders.signals.queue_order_status_notifications").start()
        self.order_notifications = patch("apps.orders.signals.send_order_notification").start()
        self.addCleanup(patch.stopall)
        self.address = CustomerAddress.objects.create(
            user=self.user,
            street_name="Backfill Street",
            city="Kampala",
            region=CustomerAddress.Region.KAMPALA_AREA,
        )

    def make_historical_paid_order(self, *, slug: str, status=Order.Status.PAID):
        order = Order.objects.create(
            tenant=self.tenant,
            user=self.user,
            address=self.address,
            slug=slug,
            status=Order.Status.PROCESSING,
            shipping_fee=Decimal("500.00"),
            discount_amount=Decimal("0.00"),
        )
        add_order_item(
            order=order,
            variant=self.variant,
            quantity=1,
            unit_price=Decimal("2000.00"),
        )
        order.recalculate_total_price()
        order.status = status
        order.save(update_fields=["status", "updated_at"])
        Payment.objects.create(
            tenant=self.tenant,
            user=self.user,
            order=order,
            provider=Payment.Provider.CASH,
            status=Payment.Status.PAID,
            currency=Payment.Currency.UGX,
            amount=order.total_price,
        )
        return order

    def test_backfill_dry_run_does_not_create_events_or_journals(self):
        self.make_historical_paid_order(slug="dry-run-backfill")
        output = StringIO()

        call_command(
            "backfill_accounting",
            "--dry-run",
            "--tenant",
            self.tenant.slug,
            stdout=output,
        )

        self.assertIn("DRY-RUN", output.getvalue())
        self.assertFalse(AccountingEvent.objects.exists())
        self.assertFalse(JournalEntry.objects.exists())

    def test_backfill_is_idempotent_and_prevents_duplicate_journals(self):
        order = self.make_historical_paid_order(slug="duplicate-backfill")

        call_command("backfill_accounting", "--tenant", self.tenant.slug, stdout=StringIO())

        self.assertEqual(
            AccountingEvent.objects.filter(
                tenant=self.tenant,
                event_type="order.paid",
                source_model="orders.Order",
                source_id=str(order.pk),
            ).count(),
            1,
        )
        self.assertEqual(
            JournalEntry.objects.filter(
                tenant=self.tenant,
                source_model="orders.Order",
                source_id=str(order.pk),
            ).count(),
            2,
        )

        second_output = StringIO()
        call_command("backfill_accounting", "--tenant", self.tenant.slug, stdout=second_output)

        self.assertIn("accounting event already exists", second_output.getvalue())
        self.assertEqual(
            AccountingEvent.objects.filter(
                tenant=self.tenant,
                event_type="order.paid",
                source_model="orders.Order",
                source_id=str(order.pk),
            ).count(),
            1,
        )
        self.assertEqual(
            JournalEntry.objects.filter(
                tenant=self.tenant,
                source_model="orders.Order",
                source_id=str(order.pk),
            ).count(),
            2,
        )

    def test_backfill_skips_missing_cost_unless_allowed(self):
        self.variant.unit_cost = Decimal("0.00")
        self.variant.save(update_fields=["unit_cost"])
        order = self.make_historical_paid_order(slug="missing-cost-backfill")
        skipped_output = StringIO()

        call_command("backfill_accounting", "--tenant", self.tenant.slug, stdout=skipped_output)

        self.assertIn("missing product cost", skipped_output.getvalue())
        self.assertFalse(
            AccountingEvent.objects.filter(
                event_type="order.paid",
                source_id=str(order.pk),
            ).exists()
        )

        call_command(
            "backfill_accounting",
            "--tenant",
            self.tenant.slug,
            "--allow-missing-cost",
            stdout=StringIO(),
        )

        event = AccountingEvent.objects.get(event_type="order.paid", source_id=str(order.pk))
        self.assertEqual(event.status, AccountingEvent.Status.PROCESSED)
        self.assertTrue(event.payload["requires_attention"])
        self.assertFalse(
            JournalEntry.objects.filter(
                tenant=self.tenant,
                source_model="orders.Order",
                source_id=str(order.pk),
                idempotency_key=f"order-paid-{order.pk}-cogs",
            ).exists()
        )
