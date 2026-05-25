from __future__ import annotations

import csv
from io import BytesIO, StringIO, TextIOWrapper
from dataclasses import dataclass
from typing import Callable

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from apps.tenants.models import Tenant

from .models import Account, AccountingEvent, InventoryMovement, JournalEntry, Refund


@dataclass(frozen=True)
class DefaultAccount:
    code: str
    name: str
    account_type: str
    normal_balance: str
    description: str = ""
    parent_code: str = ""


DEFAULT_CHART_OF_ACCOUNTS: tuple[DefaultAccount, ...] = (
    DefaultAccount(
        code="1000",
        name="Cash",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Physical cash and cash-on-delivery collections.",
    ),
    DefaultAccount(
        code="1010",
        name="Bank",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Bank accounts and card settlement balances.",
    ),
    DefaultAccount(
        code="1020",
        name="Mobile Money",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Mobile money wallet and collection balances.",
    ),
    DefaultAccount(
        code="1030",
        name="Payment Gateway Clearing",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Temporary balances due from online payment processors.",
    ),
    DefaultAccount(
        code="1040",
        name="Cash on Delivery Receivable",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="COD amounts collected by riders or agents before settlement.",
    ),
    DefaultAccount(
        code="1100",
        name="Accounts Receivable",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Amounts owed by customers.",
    ),
    DefaultAccount(
        code="1200",
        name="Inventory",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Inventory value held for sale.",
    ),
    DefaultAccount(
        code="1210",
        name="Inventory In Transit",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Purchased goods paid for but not yet received into stock.",
    ),
    DefaultAccount(
        code="1300",
        name="Prepaid Expenses",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Expenses paid in advance, including rent, software, and insurance.",
    ),
    DefaultAccount(
        code="1500",
        name="Equipment and Fixtures",
        account_type=Account.Type.ASSET,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Store, warehouse, office, and fulfillment equipment.",
    ),
    DefaultAccount(
        code="2000",
        name="Accounts Payable",
        account_type=Account.Type.LIABILITY,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Amounts owed to suppliers and vendors.",
    ),
    DefaultAccount(
        code="2100",
        name="VAT Payable",
        account_type=Account.Type.LIABILITY,
        normal_balance=Account.NormalBalance.CREDIT,
        description="VAT and sales tax collected for remittance.",
    ),
    DefaultAccount(
        code="2200",
        name="Customer Deposits",
        account_type=Account.Type.LIABILITY,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Customer advances, gift cards, wallet balances, and unfulfilled prepaid orders.",
    ),
    DefaultAccount(
        code="2300",
        name="Accrued Expenses",
        account_type=Account.Type.LIABILITY,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Expenses incurred but not yet billed or paid.",
    ),
    DefaultAccount(
        code="3000",
        name="Owner Equity",
        account_type=Account.Type.EQUITY,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Owner capital and equity contributions.",
    ),
    DefaultAccount(
        code="3100",
        name="Retained Earnings",
        account_type=Account.Type.EQUITY,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Accumulated prior-period earnings.",
    ),
    DefaultAccount(
        code="3200",
        name="Owner Drawings",
        account_type=Account.Type.EQUITY,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Owner withdrawals or distributions.",
    ),
    DefaultAccount(
        code="4000",
        name="Sales Revenue",
        account_type=Account.Type.INCOME,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Product sales revenue.",
    ),
    DefaultAccount(
        code="4010",
        name="Delivery Income",
        account_type=Account.Type.INCOME,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Delivery and shipping fees charged to customers.",
    ),
    DefaultAccount(
        code="4020",
        name="Service and Handling Income",
        account_type=Account.Type.INCOME,
        normal_balance=Account.NormalBalance.CREDIT,
        description="Handling, packaging, or service fees charged to customers.",
    ),
    DefaultAccount(
        code="4090",
        name="Discounts / Contra Revenue",
        account_type=Account.Type.INCOME,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Discounts and promotions reducing gross revenue.",
    ),
    DefaultAccount(
        code="4100",
        name="Sales Returns / Refunds",
        account_type=Account.Type.INCOME,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Returns and refunds reducing sales revenue.",
    ),
    DefaultAccount(
        code="5000",
        name="Cost of Goods Sold",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Inventory cost recognized when goods are sold.",
    ),
    DefaultAccount(
        code="5100",
        name="Delivery Expense",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Delivery, courier, and fulfillment expenses.",
    ),
    DefaultAccount(
        code="5200",
        name="Payment Processing Fees",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Gateway, card, mobile money, and processor fees.",
    ),
    DefaultAccount(
        code="5300",
        name="Packaging Supplies",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Bags, boxes, labels, and fulfillment packaging.",
    ),
    DefaultAccount(
        code="5400",
        name="Marketing and Advertising",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Digital ads, campaigns, promotions, and marketplace marketing costs.",
    ),
    DefaultAccount(
        code="5500",
        name="Platform and Software Fees",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Hosting, ecommerce platform, apps, subscriptions, and SaaS tools.",
    ),
    DefaultAccount(
        code="5600",
        name="Rent and Utilities",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Store, office, warehouse, power, water, and internet costs.",
    ),
    DefaultAccount(
        code="5700",
        name="Salaries and Wages",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Staff salaries, wages, commissions, and related labor costs.",
    ),
    DefaultAccount(
        code="5800",
        name="Bad Debt Expense",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Uncollectible customer balances and write-offs.",
    ),
    DefaultAccount(
        code="5900",
        name="General Administrative Expense",
        account_type=Account.Type.EXPENSE,
        normal_balance=Account.NormalBalance.DEBIT,
        description="Other operating and administrative expenses.",
    ),
)


DEFAULT_ACCOUNT_CODES = tuple(account.code for account in DEFAULT_CHART_OF_ACCOUNTS)
CHART_IMPORT_HEADERS = ("code", "name", "account_type", "parent_code", "description", "is_active")
CHART_OF_ACCOUNTS_TEMPLATE_ROWS = (
    ("1000", "Cash", "asset", "", "Physical cash and cash-on-delivery collections.", "true"),
    ("1010", "Bank", "asset", "", "Bank accounts and card settlement balances.", "true"),
    ("1020", "Mobile Money", "asset", "", "Mobile money wallet and collection balances.", "true"),
    ("1030", "Payment Gateway Clearing", "asset", "", "Temporary balances due from online payment processors.", "true"),
    ("1040", "Cash on Delivery Receivable", "asset", "", "COD amounts collected by riders or agents before settlement.", "true"),
    ("1100", "Accounts Receivable", "asset", "", "Amounts owed by customers.", "true"),
    ("1200", "Inventory", "asset", "", "Inventory value held for sale.", "true"),
    ("1210", "Inventory In Transit", "asset", "", "Purchased goods paid for but not yet received into stock.", "true"),
    ("1300", "Prepaid Expenses", "asset", "", "Expenses paid in advance, including rent, software, and insurance.", "true"),
    ("1500", "Equipment and Fixtures", "asset", "", "Store, warehouse, office, and fulfillment equipment.", "true"),
    ("2000", "Accounts Payable", "liability", "", "Amounts owed to suppliers and vendors.", "true"),
    ("2100", "VAT Payable", "liability", "", "VAT and sales tax collected for remittance.", "true"),
    ("2200", "Customer Deposits", "liability", "", "Customer advances, gift cards, wallet balances, and unfulfilled prepaid orders.", "true"),
    ("2300", "Accrued Expenses", "liability", "", "Expenses incurred but not yet billed or paid.", "true"),
    ("3000", "Owner Equity", "equity", "", "Owner capital and equity contributions.", "true"),
    ("3100", "Retained Earnings", "equity", "", "Accumulated prior-period earnings.", "true"),
    ("3200", "Owner Drawings", "equity", "", "Owner withdrawals or distributions.", "true"),
    ("4000", "Sales Revenue", "income", "", "Product sales revenue.", "true"),
    ("4010", "Delivery Income", "income", "", "Delivery and shipping fees charged to customers.", "true"),
    ("4020", "Service and Handling Income", "income", "", "Handling, packaging, or service fees charged to customers.", "true"),
    ("4090", "Discounts / Contra Revenue", "income", "", "Discounts and promotions reducing gross revenue.", "true"),
    ("4100", "Sales Returns / Refunds", "income", "", "Returns and refunds reducing sales revenue.", "true"),
    ("5000", "Cost of Goods Sold", "expense", "", "Inventory cost recognized when goods are sold.", "true"),
    ("5100", "Delivery Expense", "expense", "", "Delivery, courier, and fulfillment expenses.", "true"),
    ("5200", "Payment Processing Fees", "expense", "", "Gateway, card, mobile money, and processor fees.", "true"),
    ("5300", "Packaging Supplies", "expense", "", "Bags, boxes, labels, and fulfillment packaging.", "true"),
    ("5400", "Marketing and Advertising", "expense", "", "Digital ads, campaigns, promotions, and marketplace marketing costs.", "true"),
    ("5500", "Platform and Software Fees", "expense", "", "Hosting, ecommerce platform, apps, subscriptions, and SaaS tools.", "true"),
    ("5600", "Rent and Utilities", "expense", "", "Store, office, warehouse, power, water, and internet costs.", "true"),
    ("5700", "Salaries and Wages", "expense", "", "Staff salaries, wages, commissions, and related labor costs.", "true"),
    ("5800", "Bad Debt Expense", "expense", "", "Uncollectible customer balances and write-offs.", "true"),
    ("5900", "General Administrative Expense", "expense", "", "Other operating and administrative expenses.", "true"),
)


def normalize_account_code(value: object) -> str:
    return str(value or "").strip().upper()


def normalize_account_type(value: object) -> str:
    raw = str(value or "").strip().upper()
    aliases = {choice.value.lower(): choice.value for choice in Account.Type}
    aliases.update({choice.value: choice.value for choice in Account.Type})
    return aliases.get(raw.lower(), raw)


def account_type_for_export(value: str) -> str:
    return str(value or "").strip().lower()


def normal_balance_for_account_type(account_type: str) -> str:
    return (
        Account.NormalBalance.DEBIT
        if account_type in {Account.Type.ASSET, Account.Type.EXPENSE}
        else Account.NormalBalance.CREDIT
    )


def parse_bool(value: object, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "active"}


def read_chart_of_accounts_upload(uploaded_file) -> list[dict]:
    name = (getattr(uploaded_file, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip().lower() for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]

    text_stream = TextIOWrapper(uploaded_file, encoding="utf-8-sig", newline="")
    return list(csv.DictReader(text_stream))


def import_chart_of_accounts(*, tenant: Tenant, uploaded_file, update_existing: bool = False) -> dict:
    rows = read_chart_of_accounts_upload(uploaded_file)
    errors: list[dict] = []
    skipped_count = 0
    created_count = 0
    updated_count = 0
    pending: list[dict] = []
    seen_codes: set[str] = set()
    valid_types = {choice.value for choice in Account.Type}

    for row_number, row in enumerate(rows, start=2):
        code = normalize_account_code(row.get("code"))
        name = str(row.get("name") or "").strip()
        account_type = normalize_account_type(row.get("account_type"))
        parent_code = normalize_account_code(row.get("parent_code"))

        row_errors = []
        if not code:
            row_errors.append("code is required.")
        if not name:
            row_errors.append("name is required.")
        if account_type not in valid_types:
            row_errors.append("account_type must be one of asset, liability, equity, income, expense.")
        if code and code in seen_codes:
            row_errors.append("Duplicate code in uploaded file.")
        if parent_code and parent_code == code:
            row_errors.append("parent_code cannot match code.")

        if row_errors:
            errors.append({"row": row_number, "code": code, "errors": row_errors})
            skipped_count += 1
            continue

        seen_codes.add(code)
        pending.append(
            {
                "row": row_number,
                "code": code,
                "name": name,
                "account_type": account_type,
                "parent_code": parent_code,
                "description": str(row.get("description") or "").strip(),
                "is_active": parse_bool(row.get("is_active"), True),
            }
        )

    with transaction.atomic():
        accounts_by_code = {
            account.code: account
            for account in Account.objects.select_for_update().filter(tenant=tenant)
        }
        remaining = pending

        while remaining:
            next_remaining = []
            made_progress = False

            for row in remaining:
                existing = accounts_by_code.get(row["code"])
                parent = accounts_by_code.get(row["parent_code"]) if row["parent_code"] else None

                if row["parent_code"] and parent is None:
                    if any(candidate["code"] == row["parent_code"] for candidate in remaining):
                        next_remaining.append(row)
                        continue
                    errors.append(
                        {
                            "row": row["row"],
                            "code": row["code"],
                            "errors": [f"parent_code {row['parent_code']} was not found."],
                        }
                    )
                    skipped_count += 1
                    made_progress = True
                    continue

                payload = {
                    "name": row["name"],
                    "account_type": row["account_type"],
                    "normal_balance": normal_balance_for_account_type(row["account_type"]),
                    "parent": parent,
                    "description": row["description"],
                    "is_active": row["is_active"],
                    "currency": (getattr(tenant, "currency", "") or "UGX").strip().upper(),
                }

                try:
                    if existing is not None:
                        if not update_existing:
                            errors.append(
                                {
                                    "row": row["row"],
                                    "code": row["code"],
                                    "errors": ["Account code already exists for this tenant."],
                                }
                            )
                            skipped_count += 1
                            made_progress = True
                            continue
                        for key, value in payload.items():
                            setattr(existing, key, value)
                        existing.save()
                        updated_count += 1
                    else:
                        existing = Account.objects.create(tenant=tenant, code=row["code"], **payload)
                        accounts_by_code[row["code"]] = existing
                        created_count += 1
                    made_progress = True
                except ValidationError as exc:
                    errors.append(
                        {
                            "row": row["row"],
                            "code": row["code"],
                            "errors": exc.messages if hasattr(exc, "messages") else [str(exc)],
                        }
                    )
                    skipped_count += 1
                    made_progress = True

            if not made_progress:
                for row in next_remaining:
                    errors.append(
                        {
                            "row": row["row"],
                            "code": row["code"],
                            "errors": [f"parent_code {row['parent_code']} could not be resolved."],
                        }
                    )
                    skipped_count += 1
                break
            remaining = next_remaining

    return {
        "created_count": created_count,
        "updated_count": updated_count,
        "skipped_count": skipped_count,
        "errors": errors,
    }


def export_chart_of_accounts_csv(*, accounts) -> str:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(CHART_IMPORT_HEADERS)
    for account in accounts:
        writer.writerow(
            [
                account.code,
                account.name,
                account_type_for_export(account.account_type),
                account.parent.code if account.parent_id else "",
                account.description,
                "true" if account.is_active else "false",
            ]
        )
    return output.getvalue()


def export_chart_of_accounts_xlsx(*, accounts) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chart of Accounts"
    sheet.append(CHART_IMPORT_HEADERS)
    for account in accounts:
        sheet.append(
            [
                account.code,
                account.name,
                account_type_for_export(account.account_type),
                account.parent.code if account.parent_id else "",
                account.description,
                account.is_active,
            ]
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def chart_of_accounts_template_csv() -> str:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(CHART_IMPORT_HEADERS)
    writer.writerows(CHART_OF_ACCOUNTS_TEMPLATE_ROWS)
    return output.getvalue()


def chart_of_accounts_template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chart of Accounts"
    sheet.append(CHART_IMPORT_HEADERS)
    for row in CHART_OF_ACCOUNTS_TEMPLATE_ROWS:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def chart_of_accounts_pdf(*, accounts) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
        title="Chart of Accounts",
    )
    styles = getSampleStyleSheet()
    rows = [["Code", "Name", "Type", "Parent", "Description", "Active"]]
    for account in accounts:
        rows.append(
            [
                account.code,
                account.name,
                account_type_for_export(account.account_type),
                account.parent.code if account.parent_id else "",
                account.description or "",
                "Yes" if account.is_active else "No",
            ]
        )

    table = Table(rows, colWidths=[0.8 * inch, 1.6 * inch, 1.0 * inch, 0.8 * inch, 3.2 * inch, 0.7 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#127D61")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D6DEE6")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    document.build(
        [
            Paragraph("Chart of Accounts", styles["Title"]),
            Spacer(1, 0.15 * inch),
            table,
        ]
    )
    return output.getvalue()


@transaction.atomic
def seed_default_accounts_for_tenant(*, tenant: Tenant) -> dict:
    """Create missing default chart-of-accounts records for one tenant.

    Existing accounts are intentionally left untouched so operators can rename,
    deactivate, or describe accounts without the seeder undoing those choices.
    """

    currency = (getattr(tenant, "currency", "") or "UGX").strip().upper()
    created: list[Account] = []
    existing: list[Account] = []

    for default in DEFAULT_CHART_OF_ACCOUNTS:
        account, was_created = Account.objects.get_or_create(
            tenant=tenant,
            code=default.code,
            defaults={
                "name": default.name,
                "account_type": default.account_type,
                "normal_balance": default.normal_balance,
                "currency": currency,
                "description": default.description,
                "is_active": True,
            },
        )
        if was_created:
            created.append(account)
        else:
            existing.append(account)

    return {
        "tenant": tenant,
        "created": created,
        "existing": existing,
        "created_count": len(created),
        "existing_count": len(existing),
        "total_defaults": len(DEFAULT_CHART_OF_ACCOUNTS),
    }


def seed_default_accounts_for_tenants(*, tenants) -> list[dict]:
    return [
        seed_default_accounts_for_tenant(tenant=tenant)
        for tenant in tenants
    ]


AccountingEventHandler = Callable[[AccountingEvent], JournalEntry | None]
_POSTING_HANDLERS: dict[str, AccountingEventHandler] = {}


def register_accounting_event_handler(
    event_type: str,
    handler: AccountingEventHandler,
) -> None:
    normalized_event_type = event_type.strip()
    if not normalized_event_type:
        raise ValueError("event_type is required.")
    _POSTING_HANDLERS[normalized_event_type] = handler


def get_accounting_event_handler(event_type: str) -> AccountingEventHandler | None:
    return _POSTING_HANDLERS.get(event_type.strip())


def record_accounting_event(
    *,
    tenant: Tenant,
    event_type: str,
    source_model: str,
    source_id: str | int,
    idempotency_key: str,
    payload: dict | None = None,
) -> tuple[AccountingEvent, bool]:
    """Record an accounting event exactly once for its source/idempotency tuple."""

    try:
        return AccountingEvent.objects.get_or_create(
            tenant=tenant,
            event_type=event_type.strip(),
            source_model=source_model.strip(),
            source_id=str(source_id).strip(),
            idempotency_key=idempotency_key.strip(),
            defaults={"payload": payload or {}},
        )
    except IntegrityError:
        return (
            AccountingEvent.objects.get(
                tenant=tenant,
                event_type=event_type.strip(),
                source_model=source_model.strip(),
                source_id=str(source_id).strip(),
                idempotency_key=idempotency_key.strip(),
            ),
            False,
        )


def record_accounting_event_on_commit(
    *,
    tenant: Tenant,
    event_type: str,
    source_model: str,
    source_id: str | int,
    idempotency_key: str,
    payload: dict | None = None,
) -> None:
    """Schedule event recording after the current transaction commits."""

    def _record_event() -> None:
        record_accounting_event(
            tenant=tenant,
            event_type=event_type,
            source_model=source_model,
            source_id=source_id,
            idempotency_key=idempotency_key,
            payload=payload,
        )

    transaction.on_commit(_record_event)


@transaction.atomic
def process_accounting_event(*, event: AccountingEvent) -> AccountingEvent:
    """Process one accounting event through the registered posting skeleton."""

    locked = AccountingEvent.objects.select_for_update().get(pk=event.pk)
    if locked.status == AccountingEvent.Status.PROCESSED:
        return locked

    locked.status = AccountingEvent.Status.PROCESSING
    locked.error_message = ""
    locked.save(update_fields=["status", "error_message", "updated_at"])

    handler = get_accounting_event_handler(locked.event_type)
    if handler is None:
        locked.status = AccountingEvent.Status.FAILED
        locked.error_message = (
            f"No accounting posting handler registered for event type '{locked.event_type}'."
        )
        locked.processed_at = timezone.now()
        locked.save(
            update_fields=[
                "status",
                "error_message",
                "processed_at",
                "updated_at",
            ]
        )
        return locked

    try:
        journal_entry = handler(locked)
    except Exception as exc:
        locked.status = AccountingEvent.Status.FAILED
        locked.error_message = str(exc)
        locked.processed_at = timezone.now()
        locked.save(
            update_fields=[
                "status",
                "error_message",
                "processed_at",
                "updated_at",
            ]
        )
        return locked

    locked.status = AccountingEvent.Status.PROCESSED
    locked.error_message = ""
    locked.processed_at = timezone.now()
    if journal_entry is not None:
        locked.journal_entry = journal_entry
    locked.save(
        update_fields=[
            "status",
            "error_message",
            "processed_at",
            "journal_entry",
            "updated_at",
        ]
    )
    return locked


def retry_accounting_event(*, event: AccountingEvent) -> AccountingEvent:
    if event.status != AccountingEvent.Status.FAILED:
        raise ValueError("Only failed accounting events can be retried.")
    return process_accounting_event(event=event)


@transaction.atomic
def approve_refund(*, refund: Refund, user=None) -> Refund:
    locked = Refund.objects.select_for_update().get(pk=refund.pk)
    if locked.status != Refund.Status.DRAFT:
        raise ValueError("Only draft refunds can be approved.")
    locked.status = Refund.Status.APPROVED
    locked.approved_at = timezone.now()
    locked.approved_by = user if getattr(user, "is_authenticated", False) else None
    locked.save(update_fields=["status", "approved_at", "approved_by", "updated_at"])
    return locked


@transaction.atomic
def complete_refund(*, refund: Refund, user=None) -> Refund:
    from apps.accounting.posting import queue_refund_completed_accounting_event

    locked = (
        Refund.objects.select_for_update()
        .select_related("order", "payment", "tenant")
        .prefetch_related("lines", "lines__order_item", "lines__order_item__variant")
        .get(pk=refund.pk)
    )
    if locked.status != Refund.Status.APPROVED:
        raise ValueError("Only approved refunds can be completed.")

    for line in locked.lines.all():
        if not line.return_to_stock:
            continue
        variant = line.order_item.variant
        variant.stock_quantity += line.quantity
        variant.save(update_fields=["stock_quantity"])
        unit_cost = getattr(line.order_item, "cost_price_snapshot", None)
        if not unit_cost:
            unit_cost = variant.unit_cost
        InventoryMovement.objects.create(
            tenant=locked.tenant,
            variant=variant,
            movement_type=InventoryMovement.MovementType.RETURN,
            quantity=line.quantity,
            unit_cost=unit_cost,
            source_model="accounting.Refund",
            source_id=str(locked.pk),
            note=f"Refund return for order {locked.order.slug}",
            metadata={
                "refund_id": locked.pk,
                "refund_line_id": line.pk,
                "order_id": locked.order_id,
                "order_item_id": line.order_item_id,
            },
        )

    locked.status = Refund.Status.COMPLETED
    locked.completed_at = timezone.now()
    locked.completed_by = user if getattr(user, "is_authenticated", False) else None
    locked._allow_finalized_update = True
    locked.save(update_fields=["status", "completed_at", "completed_by", "updated_at"])
    queue_refund_completed_accounting_event(locked)
    return locked
